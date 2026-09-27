from __future__ import annotations

import csv
import json
import re
from pathlib import Path


def _norm_key(k):
    s = '' if k is None else str(k)
    s = s.strip().lower()
    # Make "Differential Equation", "differential_equation", etc. equivalent.
    return re.sub(r'[^a-z0-9]+', '', s)


def _pick_key(keys, candidates):
    low = {_norm_key(k): k for k in keys if k is not None and str(k).strip()}
    for c in candidates:
        nc = _norm_key(c)
        if nc in low:
            return low[nc]
    # Conservative substring fallback, useful for names like "ode_equation".
    for c in candidates:
        nc = _norm_key(c)
        if len(nc) < 4:
            continue
        for nk, original in low.items():
            if nc in nk or nk in nc:
                return original
    return None


EQ_KEYS = [
    'equation', 'differential equation', 'differential_equation', 'ode',
    'ode equation', 'input', 'problem', 'question', 'task',
]
ANS_KEYS = [
    'answer', 'solution', 'analytical solution', 'analytic solution',
    'general solution', 'target', 'output', 'result',
]
TYPE_KEYS = [
    'type', 'class', 'category', 'equation_type', 'equation type',
    'ode_type', 'ode type', 'family',
]


def _safe_cell(v):
    if v is None:
        return ''
    return str(v)


def _mapping_to_record(row, rel, source_row, sheet=''):
    keys = list(row.keys())
    eqk = _pick_key(keys, EQ_KEYS)
    ank = _pick_key(keys, ANS_KEYS)
    typek = _pick_key(keys, TYPE_KEYS)
    if not eqk or not ank:
        return None, (eqk, ank, typek)
    src = rel if not sheet else f'{rel}#{sheet}'
    return {
        'source_file': src,
        'source_row': source_row,
        'equation_latex': _safe_cell(row.get(eqk, '')),
        'answer_latex': _safe_cell(row.get(ank, '')),
        'category': _safe_cell(row.get(typek, '')) if typek else '',
    }, (eqk, ank, typek)


def _iter_xlsx(path: Path, rel: str):
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise RuntimeError(
            'Reading the HSE .xlsx files requires openpyxl. '
            'Run: pip install openpyxl>=3.1'
        ) from e

    wb = load_workbook(path, read_only=True, data_only=False)
    yielded_any = False
    try:
        for ws in wb.worksheets:
            it = ws.iter_rows(values_only=True)
            header = None
            header_excel_row = 0
            # Be tolerant to a few leading blank/title rows.
            for excel_row, values in enumerate(it, start=1):
                if values and any(v is not None and str(v).strip() for v in values):
                    header = list(values)
                    header_excel_row = excel_row
                    break
            if header is None:
                continue

            headers = []
            for i, v in enumerate(header):
                name = str(v).strip() if v is not None and str(v).strip() else f'__col_{i+1}'
                # Dicts need unique keys even if Excel has duplicate headers.
                if name in headers:
                    name = f'{name}_{i+1}'
                headers.append(name)

            eqk = _pick_key(headers, EQ_KEYS)
            ank = _pick_key(headers, ANS_KEYS)
            typek = _pick_key(headers, TYPE_KEYS)
            print(
                f'[dataset] {rel}#{ws.title}: headers={headers} '
                f'equation_col={eqk!r} answer_col={ank!r} category_col={typek!r}'
            )
            if not eqk or not ank:
                print(
                    f'[warn] skipped sheet {rel}#{ws.title}: could not identify '
                    f'equation/answer columns. Expected aliases like '
                    f'Equation/ODE and Answer/Solution.'
                )
                continue

            for excel_row, values in enumerate(it, start=header_excel_row + 1):
                vals = list(values) if values is not None else []
                if not vals or not any(v is not None and str(v).strip() for v in vals):
                    continue
                if len(vals) < len(headers):
                    vals += [None] * (len(headers) - len(vals))
                row = {headers[i]: vals[i] if i < len(vals) else None for i in range(len(headers))}
                rec, _ = _mapping_to_record(row, rel, excel_row, sheet=ws.title)
                if rec is None:
                    continue
                # Ignore entirely empty equation/answer rows.
                if not rec['equation_latex'].strip() or not rec['answer_latex'].strip():
                    continue
                yielded_any = True
                yield rec
    finally:
        wb.close()

    if not yielded_any:
        print(f'[warn] no usable equation/answer rows found in {rel}')


def iter_dataset_rows(root: str):
    """Yield HSE equation/answer rows from XLSX/CSV/JSON/JSONL files.

    The current public HSE release stores data/train.xlsx and data/test.xlsx.
    XLSX files are streamed in read-only mode so the full ~300k corpus is not
    loaded into memory.
    """
    rootp = Path(root)
    if not rootp.exists():
        raise FileNotFoundError(root)

    supported = {'.xlsx', '.csv', '.jsonl', '.json'}
    files = sorted([p for p in rootp.rglob('*') if p.is_file() and p.suffix.lower() in supported])
    if not files:
        raise RuntimeError(
            f'No XLSX/CSV/JSON/JSONL files found under {root}. '
            f'For the public HSE repo I expect data/train.xlsx and data/test.xlsx.'
        )

    for path in files:
        rel = str(path.relative_to(rootp))
        suffix = path.suffix.lower()
        try:
            if suffix == '.xlsx':
                yield from _iter_xlsx(path, rel)

            elif suffix == '.csv':
                with path.open('r', encoding='utf-8-sig', errors='replace', newline='') as f:
                    rdr = csv.DictReader(f)
                    if not rdr.fieldnames:
                        continue
                    eqk = _pick_key(rdr.fieldnames, EQ_KEYS)
                    ank = _pick_key(rdr.fieldnames, ANS_KEYS)
                    typek = _pick_key(rdr.fieldnames, TYPE_KEYS)
                    print(
                        f'[dataset] {rel}: headers={rdr.fieldnames} '
                        f'equation_col={eqk!r} answer_col={ank!r} category_col={typek!r}'
                    )
                    if not eqk or not ank:
                        continue
                    for i, row in enumerate(rdr, start=2):
                        rec, _ = _mapping_to_record(row, rel, i)
                        if rec and rec['equation_latex'].strip() and rec['answer_latex'].strip():
                            yield rec

            elif suffix == '.jsonl':
                with path.open('r', encoding='utf-8', errors='replace') as f:
                    for i, line in enumerate(f, start=1):
                        try:
                            row = json.loads(line)
                        except Exception:
                            continue
                        if not isinstance(row, dict):
                            continue
                        rec, _ = _mapping_to_record(row, rel, i)
                        if rec and rec['equation_latex'].strip() and rec['answer_latex'].strip():
                            yield rec

            else:  # .json
                with path.open('r', encoding='utf-8', errors='replace') as f:
                    obj = json.load(f)
                rows = obj if isinstance(obj, list) else obj.get('data', []) if isinstance(obj, dict) else []
                for i, row in enumerate(rows, start=1):
                    if not isinstance(row, dict):
                        continue
                    rec, _ = _mapping_to_record(row, rel, i)
                    if rec and rec['equation_latex'].strip() and rec['answer_latex'].strip():
                        yield rec

        except Exception as e:
            print(f'[warn] skipped {rel}: {type(e).__name__}: {e}')
